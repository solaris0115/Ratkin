#!/usr/bin/env python3
"""
Parse melee weapon ThingDefs from RimworldData and Project/1.6/Defs.
Extract tool-level stats (power, cooldown, damageType) and compute DPS.
Output: .cursor/def-cache/melee-weapons.md, ExelData/Melee_Weapon_Balancing.xlsx
"""
import xml.etree.ElementTree as ET
from pathlib import Path
from collections import defaultdict
from datetime import datetime

ROOT = Path(__file__).resolve().parent.parent
RIMWORLD_SOURCES = [
    ROOT / "RimworldData" / "Core" / "Defs",
    ROOT / "RimworldData" / "Royalty" / "Defs",
    ROOT / "RimworldData" / "Ideology" / "Defs",
    ROOT / "RimworldData" / "Biotech" / "Defs",
    ROOT / "RimworldData" / "Anomaly" / "Defs",
    ROOT / "RimworldData" / "Odyssey" / "Defs",
]
PROJECT_DEFS = ROOT / "Project" / "1.6" / "Defs"

# Verb classes that indicate ranged weapon (exclude these)
RANGED_VERBS = {
    "Verb_Shoot", "Verb_LaunchProjectile", "Verb_ShootOneUse",
    "Verb_BFRShoot", "Verb_RatHolicGun",
}

# Legendary quality multiplier (MeleeWeapon_DamageMultiplier)
QUALITY_LEGENDARY = 1.65

# Armor penetration: 미지정 시 power×0.015 (effective_damage 기반), 고정 지정 시 품질만
# VerbProperties.AdjustedArmorPenetration: armorPenetration<0 → damage×0.015, >=0 → ap×quality
PIERCING_DEFAULT_FACTOR = 0.015

# Material coefficients: sharp=Plasteel, blunt=Uranium (combat-coefficients.md)
# (SharpDamageMult, BluntDamageMult, CooldownMult) - cooldown lower = faster
# Default for stuff weapons: sharp→Plasteel, blunt→Uranium
MATERIAL_PLASTEEL = (1.1, 0.9, 0.8)   # sharp 기본
MATERIAL_URANIUM = (1.1, 1.5, 1.10)   # blunt 기본

# 레퍼런스용 소재별 계수 (combat-coefficients.md)
MATERIAL_REFERENCE = {
    "Steel": (1.0, 1.0, 1.0),
    "Uranium": (1.1, 1.5, 1.10),
    "Plasteel": (1.1, 0.9, 0.8),
    "Bioferrite": (1.3, 0.9, 1.0),
}

# capacity -> (damageType, armorCategory) from damage-type-mapping.md
CAPACITY_TO_ARMOR = {
    "Poke": ("Blunt", "blunt"),
    "Cut": ("Cut", "sharp"),
    "Stab": ("Stab", "sharp"),
    "Blunt": ("Blunt", "blunt"),
    "Scratch": ("Scratch", "sharp"),
    "Demolish": ("Demolish", "blunt"),
    "Bite": ("Bite", "sharp"),
    "Flame": ("Flame", "heat"),
    "Burn": ("Burn", "heat"),
    "Crush": ("Crush", "blunt"),
    "RK_HalberdCleave": ("Cut", "sharp"),
    "RK_ToolCapacity_PickaxeStab": ("Stab", "sharp"),
    "RK_ToolCapacity_ChainSword": ("Cut", "sharp"),
    "RK_ToolCapacity_MeleeExplosion": ("Bomb", "sharp"),
    "GunlanceShell_Normal": ("Bomb", "sharp"),
}


def get_text(elem, default=""):
    if elem is not None and elem.text:
        return elem.text.strip()
    return default


def get_float(val, default=0.0):
    if val is None:
        return default
    if hasattr(val, "text"):
        t = get_text(val)
    else:
        t = str(val).strip() if val else ""
    if not t:
        return default
    try:
        return float(t)
    except ValueError:
        return default


def find_direct(parent, tag):
    for c in parent:
        if c.tag.endswith("}" + tag) or c.tag == tag:
            return c
    return None


def parse_defs_from_file(path):
    """Parse all ThingDef elements from an XML file."""
    try:
        tree = ET.parse(path)
        root = tree.getroot()
    except ET.ParseError:
        return []
    ns = root.tag.split("}")[0] + "}" if "}" in root.tag else ""
    defs = []
    for def_elem in root.iter():
        if def_elem.tag.endswith("}ThingDef") or def_elem.tag == "ThingDef":
            defs.append((path, def_elem, ns))
    return defs


def get_def_name(elem, ns):
    child = find_direct(elem, "defName")
    if child is not None:
        t = get_text(child)
        if t:
            return t
    name_attr = elem.get("Name")
    if name_attr:
        return name_attr
    return None


def get_parent_name(elem, ns):
    return elem.get("ParentName")


def _resolve_parent_name(parent_name, all_defs):
    if parent_name in all_defs and parent_name != "_name_to_def":
        return parent_name
    name_to_def = all_defs.get("_name_to_def", {})
    return name_to_def.get(parent_name, parent_name)


def collect_all_thingdefs():
    """Collect all ThingDefs from all source paths."""
    all_defs = {}
    all_files = []

    for base in RIMWORLD_SOURCES:
        if base.exists():
            for p in base.rglob("*.xml"):
                all_files.append((p, "Rimworld"))
    if PROJECT_DEFS.exists():
        for p in PROJECT_DEFS.rglob("*.xml"):
            all_files.append((p, "Project"))

    name_to_def = {}

    for path, source_type in all_files:
        for _, elem, ns in parse_defs_from_file(path):
            def_name = get_def_name(elem, ns)
            if not def_name:
                continue
            name_attr = elem.get("Name")
            rel = str(path.relative_to(ROOT))
            is_abstract = elem.get("Abstract") == "True"

            if is_abstract:
                if def_name not in all_defs:
                    all_defs[def_name] = (path, elem, ns, rel, source_type)
            else:
                if def_name not in all_defs or "Project" in rel:
                    all_defs[def_name] = (path, elem, ns, rel, source_type)

            if name_attr and name_attr != def_name:
                name_to_def[name_attr] = def_name

    all_defs["_name_to_def"] = name_to_def
    return all_defs


def _has_tools(elem):
    return find_direct(elem, "tools") is not None


def _get_tools_elem(elem, all_defs, cache):
    """Get tools element from def or parent chain. Returns first found (most specific)."""
    if elem is None:
        return None
    tools = find_direct(elem, "tools")
    if tools is not None:
        return tools
    parent_name = get_parent_name(elem, None)
    if not parent_name:
        return None
    if parent_name in cache:
        return cache[parent_name]
    if parent_name not in all_defs or parent_name == "_name_to_def":
        cache[parent_name] = None
        return None
    _, parent_elem, _, _, _ = all_defs[parent_name]
    actual = _resolve_parent_name(parent_name, all_defs)
    if actual != parent_name:
        _, parent_elem, _, _, _ = all_defs.get(actual, (None, None, None, None, None))
    result = _get_tools_elem(parent_elem, all_defs, cache)
    cache[parent_name] = result
    return result


def _has_ranged_verb(elem, all_defs, cache):
    """Check if def or parent has Verb_Shoot/Verb_LaunchProjectile."""
    if elem is None:
        return False
    verbs_elem = find_direct(elem, "verbs")
    if verbs_elem is not None:
        for li in verbs_elem:
            if li.tag.endswith("}li") or li.tag == "li":
                vc = find_direct(li, "verbClass")
                if vc is not None:
                    vc_text = get_text(vc)
                    if vc_text:
                        name = vc_text.split(".")[-1] if "." in vc_text else vc_text
                        if name in RANGED_VERBS:
                            return True
    parent_name = get_parent_name(elem, None)
    if not parent_name:
        return False
    if parent_name in cache:
        return cache[parent_name]
    if parent_name not in all_defs or parent_name == "_name_to_def":
        cache[parent_name] = False
        return False
    actual = _resolve_parent_name(parent_name, all_defs)
    _, parent_elem, _, _, _ = all_defs.get(actual, (None, None, None, None, None))
    result = _has_ranged_verb(parent_elem, all_defs, cache)
    cache[parent_name] = result
    return result


def _has_stuff_categories(elem, all_defs, cache):
    """Check if def or parent has stuffCategories."""
    if elem is None:
        return False
    sc = find_direct(elem, "stuffCategories")
    if sc is not None and len(list(sc)) > 0:
        return True
    parent_name = get_parent_name(elem, None)
    if not parent_name:
        return False
    if parent_name in cache:
        return cache[parent_name]
    if parent_name not in all_defs or parent_name == "_name_to_def":
        cache[parent_name] = False
        return False
    actual = _resolve_parent_name(parent_name, all_defs)
    _, parent_elem, _, _, _ = all_defs.get(actual, (None, None, None, None, None))
    result = _has_stuff_categories(parent_elem, all_defs, cache)
    cache[parent_name] = result
    return result


def _has_cost_list(elem, all_defs, cache):
    """Check if def or parent has costList."""
    if elem is None:
        return False
    cl = find_direct(elem, "costList")
    if cl is not None and len(list(cl)) > 0:
        return True
    parent_name = get_parent_name(elem, None)
    if not parent_name:
        return False
    if parent_name in cache:
        return cache[parent_name]
    if parent_name not in all_defs or parent_name == "_name_to_def":
        cache[parent_name] = False
        return False
    actual = _resolve_parent_name(parent_name, all_defs)
    _, parent_elem, _, _, _ = all_defs.get(actual, (None, None, None, None, None))
    result = _has_cost_list(parent_elem, all_defs, cache)
    cache[parent_name] = result
    return result


def _has_recipe_maker(elem, all_defs, cache):
    """Check if def or parent has recipeMaker."""
    if elem is None:
        return False
    rm = find_direct(elem, "recipeMaker")
    if rm is not None:
        return True
    parent_name = get_parent_name(elem, None)
    if not parent_name:
        return False
    if parent_name in cache:
        return cache[parent_name]
    if parent_name not in all_defs or parent_name == "_name_to_def":
        cache[parent_name] = False
        return False
    actual = _resolve_parent_name(parent_name, all_defs)
    _, parent_elem, _, _, _ = all_defs.get(actual, (None, None, None, None, None))
    result = _has_recipe_maker(parent_elem, all_defs, cache)
    cache[parent_name] = result
    return result


def classify_production(elem, all_defs, cache):
    """Return 'stuff', 'Fixed Cost', or '생산 불가'."""
    has_stuff = _has_stuff_categories(elem, all_defs, cache)
    has_cost = _has_cost_list(elem, all_defs, cache)
    has_recipe = _has_recipe_maker(elem, all_defs, cache)

    if has_stuff:
        return "stuff"
    if has_cost:
        return "Fixed Cost"
    return "생산 불가"


def capacity_to_damage_type(capacity: str) -> tuple:
    """Return (damageType, armorCategory). Unknown capacity -> (capacity, 'sharp')."""
    return CAPACITY_TO_ARMOR.get(capacity, (capacity, "sharp"))


def format_damage_type(capacities: list, extra_damages: list) -> str:
    """Format as 'Cap(armor)' - capacity 이름으로 표시 (Blunt/Poke 구분)."""
    parts = []
    for cap in capacities:
        _, armor = capacity_to_damage_type(cap)
        parts.append(f"{cap}({armor})")
    main = "; ".join(parts)
    for ed in extra_damages:
        dmg, armor = capacity_to_damage_type(ed)
        main += f"+{dmg}({armor})"
    return main


def get_primary_armor_category(capacities: list, extra_damages: list) -> str:
    """Primary armor category for material selection. sharp > blunt > heat."""
    cats = set()
    for cap in capacities:
        _, armor = capacity_to_damage_type(cap)
        cats.add(armor)
    for ed in extra_damages:
        _, armor = capacity_to_damage_type(ed)
        cats.add(armor)
    if "sharp" in cats:
        return "sharp"
    if "blunt" in cats:
        return "blunt"
    if "heat" in cats:
        return "heat"
    return "sharp"


def extract_tools(tools_elem):
    """Extract list of tool dicts from tools element. Each: capacities, power, cooldown, extraMeleeDamages."""
    result = []
    for li in tools_elem:
        if not (li.tag.endswith("}li") or li.tag == "li"):
            continue
        capacities = []
        cap_elem = find_direct(li, "capacities")
        if cap_elem is not None:
            for cli in cap_elem:
                if cli.tag.endswith("}li") or cli.tag == "li":
                    c = get_text(cli)
                    if c:
                        capacities.append(c)

        power = get_float(find_direct(li, "power"), 0.0)
        cooldown = get_float(find_direct(li, "cooldownTime"), 0.0)
        if cooldown <= 0:
            cooldown = 2.0  # fallback
        # armorPenetration: 미지정 시 -1, 지정 시 0~1 (고정 관통력, 품질만 적용)
        armor_penetration = get_float(find_direct(li, "armorPenetration"), -1.0)

        extra_damages = []
        extra_elem = find_direct(li, "extraMeleeDamages")
        if extra_elem is not None:
            for eli in extra_elem:
                if eli.tag.endswith("}li") or eli.tag == "li":
                    edef = find_direct(eli, "def")
                    if edef is not None:
                        ed = get_text(edef)
                        if ed:
                            extra_damages.append(ed)
        # Check surpriseAttack
        surprise = find_direct(li, "surpriseAttack")
        if surprise is not None:
            se = find_direct(surprise, "extraMeleeDamages")
            if se is not None:
                for eli in se:
                    if eli.tag.endswith("}li") or eli.tag == "li":
                        edef = find_direct(eli, "def")
                        if edef is not None:
                            ed = get_text(edef)
                            if ed and ed not in extra_damages:
                                extra_damages.append(ed)

        if capacities or extra_damages:
            result.append({
                "capacities": capacities,
                "power": power,
                "cooldown": cooldown,
                "armor_penetration": armor_penetration,
                "extra_damages": extra_damages,
            })
    return result


def _effective_damage_mult(armor_cat: str, is_stuff: bool) -> float:
    """전설등급 기준 유효 피해 배율 (power × 이 값 = effective_damage)."""
    dmg_mult = QUALITY_LEGENDARY
    if is_stuff:
        if armor_cat == "sharp":
            dmg_mult *= MATERIAL_PLASTEEL[0]  # SharpDamageMultiplier
        elif armor_cat == "blunt":
            dmg_mult *= MATERIAL_URANIUM[1]  # BluntDamageMultiplier
    return dmg_mult


def calc_melee_dps(power: float, cooldown: float, armor_cat: str, is_stuff: bool) -> float:
    """DPS = (power * quality * material_dmg) / (cooldown * material_cd)."""
    if cooldown <= 0:
        return 0.0
    dmg_mult = QUALITY_LEGENDARY
    cd_mult = 1.0
    if is_stuff:
        if armor_cat == "sharp":
            dmg_mult *= MATERIAL_PLASTEEL[0]
            cd_mult = MATERIAL_PLASTEEL[2]  # CooldownMultiplier
        elif armor_cat == "blunt":
            dmg_mult *= MATERIAL_URANIUM[1]
            cd_mult = MATERIAL_URANIUM[2]
        elif armor_cat == "heat":
            cd_mult = MATERIAL_PLASTEEL[2]  # heat uses plasteel cooldown as default
    effective_damage = power * dmg_mult
    effective_cooldown = cooldown * cd_mult
    return round(effective_damage / effective_cooldown, 2)


def calc_melee_piercing(
    power: float,
    armor_penetration: float,
    armor_cat: str,
    is_stuff: bool,
) -> float:
    """관통력 계산. VerbProperties.AdjustedArmorPenetration 로직.
    - armor_penetration >= 0 (고정): ap × 품질 (소재 영향 없음)
    - armor_penetration < 0 (미지정): effective_damage × 0.015
    """
    if armor_penetration >= 0:
        return round(armor_penetration * QUALITY_LEGENDARY, 2)
    effective_damage = power * _effective_damage_mult(armor_cat, is_stuff)
    return round(effective_damage * PIERCING_DEFAULT_FACTOR, 2)


def is_melee_weapon(def_name, all_defs, tools_cache, verb_cache):
    """Pure melee: has tools, no Verb_Shoot/Verb_LaunchProjectile."""
    if def_name not in all_defs or def_name == "_name_to_def":
        return False
    path, elem, ns, _, _ = all_defs[def_name]
    tools = _get_tools_elem(elem, all_defs, tools_cache)
    if tools is None:
        return False
    if _has_ranged_verb(elem, all_defs, verb_cache):
        return False
    return True


def extract_melee_weapon_data(def_name, all_defs, tools_cache, verb_cache, prod_cache):
    """Extract tool rows for a melee weapon. Returns list of dicts (one per tool)."""
    if def_name not in all_defs:
        return []
    path, elem, ns, rel, source_type = all_defs[def_name]
    tools_elem = _get_tools_elem(elem, all_defs, tools_cache)
    if tools_elem is None:
        return []
    if _has_ranged_verb(elem, all_defs, verb_cache):
        return []

    production = classify_production(elem, all_defs, prod_cache)
    is_stuff = production == "stuff"

    tool_list = extract_tools(tools_elem)
    rows = []
    for t in tool_list:
        capacities = t["capacities"]
        extra = t["extra_damages"]
        if not capacities and not extra:
            continue
        if not capacities:
            capacities = extra[:1]  # use first extra as primary
        damage_type_str = format_damage_type(capacities, extra)
        primary_armor = get_primary_armor_category(capacities, extra)
        dps = calc_melee_dps(t["power"], t["cooldown"], primary_armor, is_stuff)
        piercing = calc_melee_piercing(
            t["power"], t["armor_penetration"], primary_armor, is_stuff
        )
        rows.append({
            "defName": def_name,
            "source": rel,
            "source_type": source_type,
            "생산": production,
            "damageType": damage_type_str,
            "primary_armor": primary_armor,
            "power": t["power"],
            "piercing": piercing,
            "armor_penetration": t["armor_penetration"],  # Def 원본값: <0=power기반, >=0=고정
            "cooldown": t["cooldown"],
            "DPS": dps,
        })
    return rows


def source_to_dlc(path):
    if "Core" in path:
        return "Core"
    if "Royalty" in path:
        return "Royalty"
    if "Ideology" in path:
        return "Ideology"
    if "Biotech" in path:
        return "Biotech"
    if "Anomaly" in path:
        return "Anomaly"
    if "Odyssey" in path:
        return "Odyssey"
    return "Other"


# Path patterns: only include defs from these (equippable melee weapons)
# Excludes: animals, mechanoids, resources (Beer, WoodLog), etc.
MELEE_WEAPON_PATH_INCLUDE = [
    "ThingDefs_Misc/Weapons/Melee",
    "ThingDefs_Misc/Weapons/Breach",
    "ThingDefs_Misc/Weapons/MeleeBladelink",
    "ThingDefs_Misc/Weapons/MeleeUltratech",
    "ThingDefs_Misc/Weapons/MeleeMedieval",
    "ThingDefs_Misc/Weapons/PsychicWeapons",
    "Odyssey/Defs/ThingDefs_Items/Items_Exotic",  # MastodonTusk, AlphaThrumboHorn
    "ThingsDefs/Weapon_",
]
MELEE_WEAPON_PATH_EXCLUDE = [
    "Races_",
    "Races_Animal",
    "Races_Mechanoid",
    "Races_Entities",
    "Races_Fleshbeasts",
    "Races_Humanlike",
    "ThingDefs_Races",  # RK_Animal 등 종족 Def
    "Items_Resource",
    "Drugs/",
    "Alcohol",
]


def _is_melee_weapon_source(path: str) -> bool:
    """Check if def source path indicates equippable melee weapon (exclude animals, resources)."""
    path_norm = path.replace("\\", "/")
    for exc in MELEE_WEAPON_PATH_EXCLUDE:
        if exc in path_norm:
            return False
    for inc in MELEE_WEAPON_PATH_INCLUDE:
        if inc in path_norm:
            return True
    return False


def source_to_ratkin_sheet(path):
    """Get Ratkin sheet name from path (Weapon_Melee, Weapon_Util, etc)."""
    path_str = path.replace("\\", "/")
    if "Weapon_Melee" in path_str:
        return "Weapon_Melee"
    if "Weapon_Util" in path_str:
        return "Weapon_Util"
    if "Weapon_DropOnly" in path_str:
        return "Weapon_DropOnly"
    if "Weapon_HighTech" in path_str:
        return "Weapon_HighTech"
    return "Ratkin_Other"


def main():
    all_defs = collect_all_thingdefs()
    tools_cache = {}
    verb_cache = {}
    prod_cache = {}

    all_rows = []
    for def_name in all_defs:
        if def_name == "_name_to_def":
            continue
        if not is_melee_weapon(def_name, all_defs, tools_cache, verb_cache):
            continue
        rows = extract_melee_weapon_data(
            def_name, all_defs, tools_cache, verb_cache, prod_cache
        )
        # Filter: only equippable melee weapons (exclude animals, resources, mechanoids)
        for r in rows:
            if _is_melee_weapon_source(r["source"]):
                all_rows.append(r)

    rimworld_rows = [r for r in all_rows if r["source_type"] == "Rimworld"]
    project_rows = [r for r in all_rows if r["source_type"] == "Project"]

    rimworld_by_dlc = defaultdict(list)
    for r in rimworld_rows:
        dlc = source_to_dlc(r["source"])
        rimworld_by_dlc[dlc].append(r)

    project_by_sheet = defaultdict(list)
    for r in project_rows:
        sheet = source_to_ratkin_sheet(r["source"])
        project_by_sheet[sheet].append(r)

    sharp_rows = [r for r in all_rows if r["primary_armor"] in ("sharp", "heat")]
    blunt_rows = [r for r in all_rows if r["primary_armor"] == "blunt"]

    return {
        "all_rows": all_rows,
        "sharp_rows": sharp_rows,
        "blunt_rows": blunt_rows,
        "rimworld_by_dlc": dict(rimworld_by_dlc),
        "project": project_rows,
        "project_by_sheet": dict(project_by_sheet),
    }


def _markdown_table(headers: list, rows: list) -> str:
    def row_str(cells):
        return "| " + " | ".join(str(c) for c in cells) + " |"
    sep = "|" + "|".join(["---"] * len(headers)) + "|"
    return "\n".join([row_str(headers), sep] + [row_str(r) for r in rows])


def _round2(val):
    """소수점 두 자리로 반올림."""
    if isinstance(val, (int, float)):
        return round(val, 2)
    return val


def _tool_row(r: dict) -> list:
    return [
        r["defName"],
        r["생산"],
        r["damageType"],
        _round2(r["power"]),
        _round2(r["piercing"]),
        _round2(r["cooldown"]),
        _round2(r["DPS"]),
    ]


def write_cache_md(result):
    out_path = ROOT / ".cursor" / "def-cache" / "melee-weapons.md"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    sources = []
    for rows in result["rimworld_by_dlc"].values():
        for r in rows:
            if r["source"] not in sources:
                sources.append(r["source"])
    for r in result["project"]:
        if r["source"] not in sources:
            sources.append(r["source"])

    today = datetime.now().strftime("%Y-%m-%d")
    yaml = f"""---
category: melee-weapons
last_updated: {today}
sources:
"""
    for s in sorted(sources)[:40]:
        yaml += f"  - {s.replace(chr(92), '/')}\n"
    yaml += """scope: 순수 근접 무기만 (원거리 무기의 근접 tools 제외)
fields: defName, 생산, damageType, power, piercing, cooldown, DPS
note: DPS/관통력 = 전설등급. 고정관통력(armorPenetration 지정)은 품질만, 미지정은 effective_damage×0.015
---
"""

    sections = [
        yaml.strip(),
        "",
        "# 근접 무기 (Melee Weapons)",
        "",
        "## damageType 매핑 및 계수",
        "",
        "- **[damage-type-mapping.md](damage-type-mapping.md)**: capacity → damageType, damageType → armorCategory",
        "- **[combat-coefficients.md](combat-coefficients.md)**: 전투 관련 계수 (등급별 피해/소재별 계수)",
        "",
        "**생산**: stuff = 소재 선택 (등급/소재 계수) | Fixed Cost = 고정 재료 (등급만) | 생산 불가",
        "",
        "## 림월드 (Core + DLC)",
        "",
    ]

    headers = ["defName", "생산", "damageType", "power", "piercing", "cooldown", "DPS"]
    dlc_order = ["Core", "Royalty", "Ideology", "Biotech", "Anomaly", "Odyssey"]
    for dlc in dlc_order:
        rows = result["rimworld_by_dlc"].get(dlc, [])
        if not rows:
            continue
        sections.append(f"### {dlc} ({len(rows)} tool행)")
        sections.append("")
        sections.append(_markdown_table(headers, [_tool_row(r) for r in rows]))
        sections.append("")
        sections.append("")

    sections.append("## 랫킨 (Ratkin)")
    sections.append("")
    sheet_order = ["Weapon_Melee", "Weapon_Util", "Weapon_DropOnly", "Weapon_HighTech", "Ratkin_Other"]
    for sheet in sheet_order:
        rows = result["project_by_sheet"].get(sheet, [])
        if not rows:
            continue
        sections.append(f"### {sheet} ({len(rows)} tool행)")
        sections.append("")
        sections.append(_markdown_table(headers, [_tool_row(r) for r in rows]))
        sections.append("")
        sections.append("")

    out_path.write_text("\n".join(sections), encoding="utf-8")
    print(f"Wrote {out_path}")


def _dps_formula(row_idx: int, params_sheet: str = "Parameters") -> str:
    """DPS 수식: power(D), cooldown(F), 생산(B), damageType(C) 참조. 계수 시트 레퍼런스."""
    # 열: A=defName, B=생산, C=damageType, D=power, E=piercing, F=cooldown, G=DPS
    # Parameters: B2=QUALITY, B3=PLASTEEL_DMG, B4=PLASTEEL_CD, B5=URANIUM_DMG, B6=URANIUM_CD
    p = params_sheet
    b, c, d, f = f"B{row_idx}", f"C{row_idx}", f"D{row_idx}", f"F{row_idx}"
    return (
        f"=ROUND(({d}*IF(AND({b}=\"stuff\",ISNUMBER(SEARCH(\"(sharp)\",{c}))),{p}!$B$2*{p}!$B$3,"
        f"IF(AND({b}=\"stuff\",ISNUMBER(SEARCH(\"(blunt)\",{c}))),{p}!$B$2*{p}!$B$5,{p}!$B$2)))/"
        f"({f}*IF(AND({b}=\"stuff\",OR(ISNUMBER(SEARCH(\"(sharp)\",{c})),ISNUMBER(SEARCH(\"(heat)\",{c})))),{p}!$B$4,"
        f"IF(AND({b}=\"stuff\",ISNUMBER(SEARCH(\"(blunt)\",{c}))),{p}!$B$6,1))),2)"
    )


def _piercing_formula(row_idx: int, params_sheet: str = "Parameters") -> str:
    """관통력 수식: power 기반(armorPenetration 미지정) 시 power(D)×quality×소재계수×0.015.
    Edit_Values Slot1(sharp/heat), Slot2(blunt) 레퍼런스. power 수정 시 자동 반영."""
    # 열: A=defName, B=생산, C=damageType, D=power, E=piercing
    # Parameters: B2=QUALITY, B7=PIERCING_FACTOR(0.015)
    # Edit_Values: B2=Plasteel(sharp), B4=Uranium(blunt) - 소재별 DMG_MULT
    p = params_sheet
    ev = "Edit_Values"
    b, c, d = f"B{row_idx}", f"C{row_idx}", f"D{row_idx}"
    # effective_damage = power × quality × material_dmg_mult; piercing = effective_damage × 0.015
    mat_mult = (
        f"IF(AND({b}=\"stuff\",OR(ISNUMBER(SEARCH(\"(sharp)\",{c})),ISNUMBER(SEARCH(\"(heat)\",{c})))),{ev}!$B$2,"
        f"IF(AND({b}=\"stuff\",ISNUMBER(SEARCH(\"(blunt)\",{c}))),{ev}!$B$4,1))"
    )
    return f"=ROUND({d}*{p}!$B$2*{mat_mult}*{p}!$B$7,2)"


def write_excel(result):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    out_path = ROOT / "ExelData" / "Melee_Weapon_Balancing.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    num_fmt = "0.00"

    headers = ["defName", "생산", "damageType", "power", "piercing", "cooldown", "DPS"]

    # 1. Edit_Values 시트: 사용자가 직접 수정. Parameters가 여기 참조.
    ws_edit = wb.active
    ws_edit.title = "Edit_Values"
    ws_edit.cell(row=1, column=1, value="Slot 1 (sharp)")
    ws_edit.cell(row=1, column=2, value="DMG_MULT")
    ws_edit.cell(row=1, column=3, value="CD_MULT")
    ws_edit.cell(row=2, column=1, value="기본: Plasteel")
    ws_edit.cell(row=2, column=2, value=MATERIAL_PLASTEEL[0])
    ws_edit.cell(row=2, column=3, value=MATERIAL_PLASTEEL[2])
    ws_edit.cell(row=3, column=1, value="Slot 2 (blunt)")
    ws_edit.cell(row=3, column=2, value="DMG_MULT")
    ws_edit.cell(row=3, column=3, value="CD_MULT")
    ws_edit.cell(row=4, column=1, value="기본: Uranium")
    ws_edit.cell(row=4, column=2, value=MATERIAL_URANIUM[1])
    ws_edit.cell(row=4, column=3, value=MATERIAL_URANIUM[2])
    for r in (2, 4):
        for c in (2, 3):
            ws_edit.cell(row=r, column=c).number_format = num_fmt
    ws_edit.column_dimensions["A"].width = 18
    ws_edit.column_dimensions["B"].width = 12
    ws_edit.column_dimensions["C"].width = 12

    # 참조란: 소재별 sharp, blunt, cooldown (수정 시 참고). 관통력도 Slot1/Slot2 DMG_MULT 참조
    ws_edit.cell(row=6, column=1, value="참조: 소재별 계수 (직접 수정 시 참고). piercing=power×quality×DMG_MULT×0.015")
    ws_edit.cell(row=6, column=1).font = Font(bold=True)
    ref_headers = ["소재", "Sharp", "Blunt", "Cooldown"]
    for c, h in enumerate(ref_headers, 1):
        cell = ws_edit.cell(row=7, column=c, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
    for i, (mat_name, (sharp, blunt, cd)) in enumerate(MATERIAL_REFERENCE.items(), 8):
        ws_edit.cell(row=i, column=1, value=mat_name)
        ws_edit.cell(row=i, column=2, value=sharp)
        ws_edit.cell(row=i, column=3, value=blunt)
        ws_edit.cell(row=i, column=4, value=cd)
        for c in (2, 3, 4):
            ws_edit.cell(row=i, column=c).number_format = num_fmt

    # 2. Material_Reference 시트: 레퍼런스용 참고 데이터만
    ws_ref = wb.create_sheet("Material_Reference")
    ws_ref.cell(row=1, column=1, value="소재")
    ws_ref.cell(row=1, column=2, value="SharpDamageMult")
    ws_ref.cell(row=1, column=3, value="BluntDamageMult")
    ws_ref.cell(row=1, column=4, value="CooldownMult")
    ws_ref.cell(row=2, column=1, value="참조: combat-coefficients.md")
    for c in range(1, 5):
        ws_ref.cell(row=1, column=c).font = header_font_white
        ws_ref.cell(row=1, column=c).fill = header_fill
    for i, (mat_name, (sharp, blunt, cd)) in enumerate(MATERIAL_REFERENCE.items(), 3):
        ws_ref.cell(row=i, column=1, value=mat_name)
        ws_ref.cell(row=i, column=2, value=sharp)
        ws_ref.cell(row=i, column=3, value=blunt)
        ws_ref.cell(row=i, column=4, value=cd)
        for c in (2, 3, 4):
            ws_ref.cell(row=i, column=c).number_format = num_fmt
    ws_ref.column_dimensions["A"].width = 14
    ws_ref.column_dimensions["B"].width = 16
    ws_ref.column_dimensions["C"].width = 16
    ws_ref.column_dimensions["D"].width = 14

    # 3. Parameters 시트: Edit_Values 링크 참조
    ws_params = wb.create_sheet("Parameters")
    ws_params.cell(row=1, column=1, value="계수")
    ws_params.cell(row=1, column=2, value="값")
    ws_params.cell(row=1, column=3, value="설명")
    for c in range(1, 4):
        ws_params.cell(row=1, column=c).font = header_font_white
        ws_params.cell(row=1, column=c).fill = header_fill
    ws_params.cell(row=2, column=1, value="QUALITY_LEGENDARY")
    ws_params.cell(row=2, column=2, value=QUALITY_LEGENDARY)
    ws_params.cell(row=2, column=2).number_format = num_fmt
    ws_params.cell(row=2, column=3, value="전설등급 피해 배율")
    ws_params.cell(row=3, column=1, value="PLASTEEL_DMG")
    ws_params.cell(row=3, column=2, value="=Edit_Values!$B$2")
    ws_params.cell(row=3, column=3, value="sharp 소재 피해 배율 (Edit_Values Slot1)")
    ws_params.cell(row=4, column=1, value="PLASTEEL_CD")
    ws_params.cell(row=4, column=2, value="=Edit_Values!$C$2")
    ws_params.cell(row=4, column=3, value="sharp 소재 쿨다운 배율 (Edit_Values Slot1)")
    ws_params.cell(row=5, column=1, value="URANIUM_DMG")
    ws_params.cell(row=5, column=2, value="=Edit_Values!$B$4")
    ws_params.cell(row=5, column=3, value="blunt 소재 피해 배율 (Edit_Values Slot2)")
    ws_params.cell(row=6, column=1, value="URANIUM_CD")
    ws_params.cell(row=6, column=2, value="=Edit_Values!$C$4")
    ws_params.cell(row=6, column=3, value="blunt 소재 쿨다운 배율 (Edit_Values Slot2)")
    ws_params.cell(row=7, column=1, value="PIERCING_FACTOR")
    ws_params.cell(row=7, column=2, value=PIERCING_DEFAULT_FACTOR)
    ws_params.cell(row=7, column=2).number_format = "0.00"
    ws_params.cell(row=7, column=3, value="관통력 power기반 계산: effective_damage×이값 (armorPenetration 미지정 시)")
    ws_params.column_dimensions["A"].width = 20
    ws_params.column_dimensions["B"].width = 18
    ws_params.column_dimensions["C"].width = 35

    def _write_sheet(ws, title, rows, ref_info=None):
        ws.title = title[:31]
        row_offset = 1
        if ref_info:
            ref_cell = ws.cell(row=1, column=1, value=ref_info)
            ref_cell.hyperlink = "#Edit_Values!A1"
            ref_cell.font = Font(color="0563C1", underline="single")
            row_offset = 2
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_offset, column=col_idx, value=h)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for row_idx, r in enumerate(rows, row_offset + 1):
            vals = _tool_row(r)
            for col_idx, val in enumerate(vals, 1):
                h = headers[col_idx - 1]
                if h == "DPS":
                    cell = ws.cell(row=row_idx, column=col_idx, value=_dps_formula(row_idx))
                elif h == "piercing" and r.get("armor_penetration", -1) < 0:
                    # power 기반 관통: power 수정 시 자동 반영, 소재(Edit_Values) 레퍼런스
                    cell = ws.cell(row=row_idx, column=col_idx, value=_piercing_formula(row_idx))
                else:
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                if h in ("power", "piercing", "cooldown") and isinstance(val, (int, float)):
                    cell.number_format = num_fmt
                    cell.alignment = Alignment(horizontal="right")
                elif h == "DPS" or (h == "piercing" and r.get("armor_penetration", -1) < 0):
                    cell.number_format = num_fmt
                    cell.alignment = Alignment(horizontal="right")

        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(headers[col_idx - 1])
            for row_idx in range(row_offset + 1, len(rows) + row_offset + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = min(max_len + 3, 40)

        if rows:
            ws.auto_filter.ref = f"A{row_offset}:{get_column_letter(len(headers))}{len(rows) + row_offset}"

    # 4. Sharp 시트 (기본 소재: Plasteel)
    _write_sheet(wb.create_sheet(), "Sharp", result["sharp_rows"],
                 ref_info="참조: Edit_Values Slot1 (Plasteel) - 클릭하여 계수 수정")

    # 5. Blunt 시트 (기본 소재: Uranium)
    _write_sheet(wb.create_sheet(), "Blunt", result["blunt_rows"],
                 ref_info="참조: Edit_Values Slot2 (Uranium) - 클릭하여 계수 수정")

    all_rimworld = []
    for rows in result["rimworld_by_dlc"].values():
        all_rimworld.extend(rows)

    _write_sheet(wb.create_sheet(), "Rimworld", all_rimworld)
    _write_sheet(wb.create_sheet(), "Ratkin", result["project"])

    dlc_order = ["Core", "Royalty", "Ideology", "Biotech", "Anomaly", "Odyssey"]
    for dlc in dlc_order:
        rows = result["rimworld_by_dlc"].get(dlc, [])
        if rows:
            ws_dlc = wb.create_sheet()
            _write_sheet(ws_dlc, dlc, rows)

    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    result = main()
    write_cache_md(result)
    write_excel(result)
